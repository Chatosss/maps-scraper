#!/usr/bin/env python3
"""
Google Maps Scraper — Interface Gráfica
Uso: python app.py
"""

import asyncio
import csv
import queue
import threading
from datetime import datetime
from pathlib import Path
from tkinter import filedialog, messagebox, ttk
from typing import Dict, List, Optional

import customtkinter as ctk
from playwright.async_api import async_playwright

from maps_scraper import (
    MAX_IG_WORKERS,
    collect_urls,
    dismiss_consent,
    extract_business,
    find_instagram_on_site,
)
import database

# ── Tema ──────────────────────────────────────────────────────────────────────
ctk.set_appearance_mode("dark")
ctk.set_default_color_theme("blue")

SIDEBAR_W  = 270
CLR_ACCENT = "#3b82f6"
CLR_GREEN  = "#22c55e"
CLR_RED    = "#ef4444"
CLR_MUTED  = "#64748b"
CLR_BG2    = "#1e293b"
CLR_BG3    = "#0f172a"
# ─────────────────────────────────────────────────────────────────────────────

CSV_KEYS    = ["nome","categoria","avaliacao","total_avaliacoes","telefone","website","instagram","endereco","maps_url"]
CSV_HEADERS = ["Nome","Categoria","Avaliação","Avaliações","Telefone","Website","Instagram","Endereço","Maps URL"]

TABLE_COLS = [
    ("nome",       "Nome",       220, "w"),
    ("categoria",  "Categoria",  130, "center"),
    ("telefone",   "Telefone",   120, "center"),
    ("website",    "Website",    170, "w"),
    ("instagram",  "Instagram",  150, "w"),
    ("avaliacao",  "Avaliação",   70, "center"),
    ("endereco",   "Endereço",   220, "w"),
]


class App(ctk.CTk):
    def __init__(self):
        super().__init__()

        self.title("Google Maps Scraper")
        self.geometry("1280x780")
        self.minsize(1050, 650)
        self.configure(fg_color=CLR_BG3)

        self._results: List[Dict]            = []
        self._thread:  Optional[threading.Thread] = None
        self._stop   = threading.Event()
        self._queue: queue.Queue             = queue.Queue()

        database.init_db()

        self._build()
        self._poll()
        self._refresh_db_stats()

    # ── Layout ────────────────────────────────────────────────────────────────

    def _build(self):
        self.grid_columnconfigure(1, weight=1)
        self.grid_rowconfigure(0, weight=1)

        self._build_sidebar()
        self._build_main()

    def _build_sidebar(self):
        sb = ctk.CTkFrame(self, width=SIDEBAR_W, corner_radius=0, fg_color=CLR_BG2)
        sb.grid(row=0, column=0, sticky="nsew")
        sb.grid_propagate(False)
        sb.grid_columnconfigure(0, weight=1)

        row = 0

        # Header
        ctk.CTkLabel(
            sb, text="Maps Scraper",
            font=ctk.CTkFont(size=20, weight="bold"),
            text_color="white",
        ).grid(row=row, column=0, padx=20, pady=(28, 2), sticky="w")
        row += 1

        ctk.CTkLabel(
            sb, text="Extrai dados de empresas\ndo Google Maps",
            font=ctk.CTkFont(size=11),
            text_color=CLR_MUTED,
            justify="left",
        ).grid(row=row, column=0, padx=20, pady=(0, 24), sticky="w")
        row += 1

        # ── Campos ────────────────────────────────────────────────────────
        for label, placeholder, attr in [
            ("Nicho / Segmento",  'ex: "clínicas estética"', "_q_entry"),
            ("Região / Cidade",   'ex: "São Paulo, SP"',     "_l_entry"),
        ]:
            ctk.CTkLabel(
                sb, text=label,
                font=ctk.CTkFont(size=12, weight="bold"),
            ).grid(row=row, column=0, padx=20, pady=(0, 4), sticky="w")
            row += 1

            entry = ctk.CTkEntry(
                sb, placeholder_text=placeholder, height=38,
                border_color=CLR_MUTED, border_width=1,
            )
            entry.grid(row=row, column=0, padx=20, pady=(0, 16), sticky="ew")
            setattr(self, attr, entry)
            row += 1

        # ── Opções ────────────────────────────────────────────────────────
        ctk.CTkLabel(
            sb, text="Opções",
            font=ctk.CTkFont(size=12, weight="bold"),
        ).grid(row=row, column=0, padx=20, pady=(0, 8), sticky="w")
        row += 1

        self._opt_ig = ctk.CTkCheckBox(
            sb, text="Buscar Instagram nos sites",
            font=ctk.CTkFont(size=12),
        )
        self._opt_ig.select()
        self._opt_ig.grid(row=row, column=0, padx=20, pady=(0, 8), sticky="w")
        row += 1

        self._opt_headless = ctk.CTkCheckBox(
            sb, text="Headless (sem janela browser)",
            font=ctk.CTkFont(size=12),
        )
        self._opt_headless.grid(row=row, column=0, padx=20, pady=(0, 8), sticky="w")
        row += 1

        self._opt_force = ctk.CTkCheckBox(
            sb, text="Forçar re-scrape (ignorar cache)",
            font=ctk.CTkFont(size=12),
            text_color=CLR_MUTED,
        )
        self._opt_force.grid(row=row, column=0, padx=20, pady=(0, 20), sticky="w")
        row += 1

        # ── Botão iniciar ─────────────────────────────────────────────────
        self._btn_start = ctk.CTkButton(
            sb, text="▶  Iniciar Busca",
            height=44, font=ctk.CTkFont(size=14, weight="bold"),
            fg_color=CLR_ACCENT, hover_color="#2563eb",
            command=self._start,
        )
        self._btn_start.grid(row=row, column=0, padx=20, pady=(0, 8), sticky="ew")
        row += 1

        self._btn_stop = ctk.CTkButton(
            sb, text="■  Parar",
            height=36, font=ctk.CTkFont(size=13),
            fg_color=CLR_RED, hover_color="#dc2626",
            command=self._stop_scrape,
        )
        self._btn_stop.grid(row=row, column=0, padx=20, pady=(0, 0), sticky="ew")
        self._btn_stop.grid_remove()
        row += 1

        # Spacer
        sb.grid_rowconfigure(row, weight=1)
        row += 1

        # ── Exportar ──────────────────────────────────────────────────────
        ctk.CTkFrame(sb, height=1, fg_color=CLR_MUTED).grid(
            row=row, column=0, padx=20, pady=(0, 14), sticky="ew"
        )
        row += 1

        ctk.CTkLabel(
            sb, text="Exportar resultados",
            font=ctk.CTkFont(size=12, weight="bold"),
        ).grid(row=row, column=0, padx=20, pady=(0, 10), sticky="w")
        row += 1

        self._btn_xlsx = ctk.CTkButton(
            sb, text="📊  Excel (.xlsx)",
            height=38, font=ctk.CTkFont(size=12),
            fg_color="#374151", hover_color="#16a34a",
            text_color="#6b7280",
            command=self._export_xlsx, state="disabled",
        )
        self._btn_xlsx.grid(row=row, column=0, padx=20, pady=(0, 8), sticky="ew")
        row += 1

        self._btn_csv = ctk.CTkButton(
            sb, text="📄  CSV",
            height=38, font=ctk.CTkFont(size=12),
            fg_color="#374151", hover_color="#0369a1",
            text_color="#6b7280",
            command=self._export_csv, state="disabled",
        )
        self._btn_csv.grid(row=row, column=0, padx=20, pady=(0, 12), sticky="ew")
        row += 1

        # ── Banco de dados ─────────────────────────────────────────────────
        ctk.CTkFrame(sb, height=1, fg_color=CLR_MUTED).grid(
            row=row, column=0, padx=20, pady=(0, 10), sticky="ew"
        )
        row += 1

        db_header = ctk.CTkFrame(sb, fg_color="transparent")
        db_header.grid(row=row, column=0, padx=20, pady=(0, 4), sticky="ew")
        db_header.grid_columnconfigure(0, weight=1)
        ctk.CTkLabel(
            db_header, text="Banco de Dados",
            font=ctk.CTkFont(size=12, weight="bold"),
            anchor="w",
        ).grid(row=0, column=0, sticky="w")
        ctk.CTkButton(
            db_header, text="Histórico",
            width=70, height=22, font=ctk.CTkFont(size=11),
            fg_color="#334155", hover_color="#475569",
            command=self._show_history,
        ).grid(row=0, column=1, sticky="e")
        row += 1

        self._lbl_db_stats = ctk.CTkLabel(
            sb, text="Carregando…",
            font=ctk.CTkFont(size=11),
            text_color=CLR_MUTED,
            justify="left",
            anchor="w",
        )
        self._lbl_db_stats.grid(row=row, column=0, padx=20, pady=(0, 20), sticky="w")

    def _build_main(self):
        main = ctk.CTkFrame(self, corner_radius=0, fg_color="transparent")
        main.grid(row=0, column=1, sticky="nsew")
        main.grid_columnconfigure(0, weight=1)
        main.grid_rowconfigure(1, weight=1)

        # ── Status bar ────────────────────────────────────────────────────
        status = ctk.CTkFrame(main, height=72, corner_radius=10, fg_color=CLR_BG2)
        status.grid(row=0, column=0, padx=16, pady=(16, 8), sticky="ew")
        status.grid_columnconfigure(0, weight=1)
        status.grid_propagate(False)

        self._lbl_status = ctk.CTkLabel(
            status, text="Pronto. Configure a busca e clique em Iniciar.",
            font=ctk.CTkFont(size=13),
            anchor="w",
        )
        self._lbl_status.grid(row=0, column=0, padx=16, pady=(10, 2), sticky="ew")

        self._progress = ctk.CTkProgressBar(status, height=7, progress_color=CLR_ACCENT)
        self._progress.set(0)
        self._progress.grid(row=1, column=0, padx=16, pady=(0, 10), sticky="ew")

        self._lbl_count = ctk.CTkLabel(
            status, text="",
            font=ctk.CTkFont(size=11),
            text_color=CLR_MUTED,
        )
        self._lbl_count.grid(row=0, column=1, padx=(4, 16), pady=(10, 2))

        # ── Tabela ────────────────────────────────────────────────────────
        table_frame = ctk.CTkFrame(main, corner_radius=10, fg_color=CLR_BG2)
        table_frame.grid(row=1, column=0, padx=16, pady=(0, 16), sticky="nsew")
        table_frame.grid_rowconfigure(0, weight=1)
        table_frame.grid_columnconfigure(0, weight=1)

        self._build_table(table_frame)

    def _build_table(self, parent):
        style = ttk.Style()
        style.theme_use("clam")
        style.configure(
            "Maps.Treeview",
            background="#1e293b",
            foreground="#e2e8f0",
            rowheight=30,
            fieldbackground="#1e293b",
            font=("Segoe UI", 10),
            borderwidth=0,
        )
        style.configure(
            "Maps.Treeview.Heading",
            background="#0f172a",
            foreground="#93c5fd",
            font=("Segoe UI", 10, "bold"),
            relief="flat",
            padding=(8, 6),
        )
        style.map(
            "Maps.Treeview",
            background=[("selected", "#1d4ed8")],
            foreground=[("selected", "white")],
        )
        style.configure("Vertical.TScrollbar", background="#334155", troughcolor="#1e293b")
        style.configure("Horizontal.TScrollbar", background="#334155", troughcolor="#1e293b")

        cols = [c[0] for c in TABLE_COLS]
        self._tree = ttk.Treeview(
            parent, columns=cols, show="headings",
            style="Maps.Treeview", selectmode="browse",
        )

        for key, label, width, anchor in TABLE_COLS:
            self._tree.heading(key, text=label)
            self._tree.column(key, width=width, minwidth=60, anchor=anchor)

        vsb = ttk.Scrollbar(parent, orient="vertical",   command=self._tree.yview)
        hsb = ttk.Scrollbar(parent, orient="horizontal", command=self._tree.xview)
        self._tree.configure(yscrollcommand=vsb.set, xscrollcommand=hsb.set)

        self._tree.grid(row=0, column=0, sticky="nsew", padx=(8, 0), pady=8)
        vsb.grid(row=0, column=1, sticky="ns",  pady=8)
        hsb.grid(row=1, column=0, sticky="ew",  padx=(8, 0))

        # Linhas alternadas
        self._tree.tag_configure("odd",    background="#1e293b")
        self._tree.tag_configure("even",   background="#253047")
        self._tree.tag_configure("cached", background="#1a2e1a", foreground="#86efac")
        self._tree.tag_configure("ig",     foreground="#f472b6")

    # ── Polling da fila ───────────────────────────────────────────────────────

    def _poll(self):
        try:
            while True:
                self._dispatch(self._queue.get_nowait())
        except queue.Empty:
            pass
        self.after(80, self._poll)

    def _dispatch(self, msg: dict):
        t = msg["type"]

        if t == "status":
            self._lbl_status.configure(text=msg["text"])

        elif t == "phase1_pulse":
            # Fase 1: progresso indeterminado — barra fica a 15%
            self._progress.set(0.15)
            self._lbl_count.configure(text="buscando…")

        elif t == "phase1_count":
            count = msg["count"]
            self._lbl_count.configure(text=f"{count} URLs")

        elif t == "progress":
            cur, tot = msg["current"], msg["total"]
            self._progress.set(cur / tot if tot else 0)
            self._lbl_count.configure(text=f"{cur} / {tot}")

        elif t == "result":
            data       = msg["data"]
            idx        = msg["idx"]
            from_cache = msg.get("from_cache", False)
            self._results.append(data)
            tag = "cached" if from_cache else ("even" if idx % 2 == 0 else "odd")
            self._tree.insert(
                "", "end", iid=str(idx), tags=(tag,),
                values=tuple(data.get(k, "") for k, *_ in TABLE_COLS),
            )
            self._lbl_count.configure(text=f"{len(self._results)} empresas")

        elif t == "ig_update":
            idx = msg["idx"]
            ig  = msg["instagram"]
            self._results[idx]["instagram"] = ig
            vals = list(self._tree.item(str(idx))["values"])
            # índice do instagram na tabela
            ig_col = [c[0] for c in TABLE_COLS].index("instagram")
            vals[ig_col] = ig
            current_tags = list(self._tree.item(str(idx))["tags"])
            if "ig" not in current_tags:
                current_tags.append("ig")
            self._tree.item(str(idx), values=vals, tags=current_tags)

        elif t == "done":
            n = len(self._results)
            ig = sum(1 for r in self._results if r.get("instagram"))
            self._lbl_status.configure(
                text=f"✓ Concluído  —  {n} empresas  |  "
                     f"{sum(1 for r in self._results if r.get('website'))} sites  |  "
                     f"{ig} instagrams  |  "
                     f"{sum(1 for r in self._results if r.get('telefone'))} telefones"
            )
            self._progress.set(1)
            self._set_running(False)
            self._refresh_db_stats()
            if self._results:
                self._btn_xlsx.configure(
                    state="normal", fg_color="#16a34a",
                    hover_color=CLR_GREEN, text_color="white",
                )
                self._btn_csv.configure(
                    state="normal", fg_color="#0369a1",
                    hover_color="#0284c7", text_color="white",
                )

        elif t == "error":
            self._lbl_status.configure(text=f"⚠  {msg['text']}")
            self._set_running(False)

    # ── Controle de estado ────────────────────────────────────────────────────

    def _set_running(self, running: bool):
        if running:
            self._btn_start.configure(state="disabled", text="Buscando…")
            self._btn_stop.grid()
        else:
            self._btn_start.configure(state="normal", text="▶  Iniciar Busca")
            self._btn_stop.configure(state="normal", text="■  Parar")
            self._btn_stop.grid_remove()

    # ── Iniciar / parar ───────────────────────────────────────────────────────

    def _start(self):
        if self._thread and self._thread.is_alive():
            return

        query    = self._q_entry.get().strip()
        location = self._l_entry.get().strip()

        if not query:
            messagebox.showwarning("Atenção", "Digite o nicho para buscar.")
            return

        self._results.clear()
        for item in self._tree.get_children():
            self._tree.delete(item)
        self._stop.clear()
        self._progress.set(0)
        self._lbl_count.configure(text="")
        self._btn_xlsx.configure(state="disabled", fg_color="#374151", text_color="#6b7280")
        self._btn_csv.configure(state="disabled",  fg_color="#374151", text_color="#6b7280")
        self._set_running(True)

        headless      = bool(self._opt_headless.get())
        scan_ig       = bool(self._opt_ig.get())
        force_refresh = bool(self._opt_force.get())

        self._thread = threading.Thread(
            target=lambda: asyncio.run(
                self._scrape_async(query, location, headless, scan_ig, force_refresh)
            ),
            daemon=True,
        )
        self._thread.start()

    def _stop_scrape(self):
        self._stop.set()
        self._btn_stop.configure(state="disabled", text="Parando…")

    # ── Lógica de scraping (thread de background) ─────────────────────────────

    async def _scrape_async(
        self, query: str, location: str,
        headless: bool, scan_ig: bool, force_refresh: bool,
    ):
        q = self._queue
        full_query = f"{query} {location}".strip() if location else query

        try:
            async with async_playwright() as pw:
                browser = await pw.chromium.launch(
                    headless=headless,
                    args=["--lang=pt-BR,pt", "--disable-blink-features=AutomationControlled"],
                )
                ctx = await browser.new_context(
                    locale="pt-BR",
                    timezone_id="America/Sao_Paulo",
                    viewport={"width": 1366, "height": 900},
                    user_agent=(
                        "Mozilla/5.0 (Windows NT 10.0; Win64; x64) "
                        "AppleWebKit/537.36 (KHTML, like Gecko) "
                        "Chrome/125.0.0.0 Safari/537.36"
                    ),
                )

                # ── Fase 1: Coletar URLs ───────────────────────────────────
                q.put({"type": "status", "text": "Fase 1  —  Coletando URLs das empresas…"})
                q.put({"type": "phase1_pulse"})

                def url_update(added, total):
                    q.put({"type": "status",
                           "text": f"Fase 1  —  {total} empresas encontradas…"})
                    q.put({"type": "phase1_count", "count": total})

                search_page = await ctx.new_page()
                urls = await collect_urls(search_page, full_query, on_update=url_update)
                await search_page.close()

                if not urls:
                    q.put({"type": "error", "text": "Nenhuma empresa encontrada para esta busca."})
                    await browser.close()
                    return

                # ── Fase 2: Extrair / carregar do cache ────────────────────
                total       = len(urls)
                new_count   = 0
                cached_count = 0

                q.put({"type": "status",   "text": f"Fase 2  —  Verificando {total} empresas…"})
                q.put({"type": "progress", "current": 0, "total": total})

                local: List[tuple] = []   # (idx, data) — para a Fase 3

                detail_page = await ctx.new_page()
                for i, url in enumerate(urls):
                    if self._stop.is_set():
                        break

                    q.put({"type": "progress", "current": i + 1, "total": total})

                    # Tentar carregar do banco primeiro
                    if not force_refresh:
                        cached = database.get_business(url)
                        if cached:
                            idx = len(local)
                            local.append((idx, cached))
                            q.put({"type": "result", "idx": idx,
                                   "data": cached, "from_cache": True})
                            cached_count += 1
                            continue

                    # Scrape fresco
                    data = await extract_business(detail_page, url)
                    if data:
                        database.save_business(data)   # salvar no banco
                        idx = len(local)
                        local.append((idx, data))
                        q.put({"type": "result", "idx": idx,
                               "data": data, "from_cache": False})
                        new_count += 1

                await detail_page.close()

                q.put({"type": "status",
                       "text": f"Fase 2 concluída  —  "
                               f"{new_count} novas  +  {cached_count} do cache"})

                # ── Fase 3: Instagram ──────────────────────────────────────
                if scan_ig and not self._stop.is_set():
                    to_scan = [
                        (idx, d) for idx, d in local
                        if d.get("website") and not d.get("instagram")
                    ]

                    if to_scan:
                        q.put({"type": "status",
                               "text": f"Fase 3  —  Buscando Instagram em {len(to_scan)} sites…"})
                        q.put({"type": "progress", "current": 0, "total": len(to_scan)})

                        sem        = asyncio.Semaphore(MAX_IG_WORKERS)
                        done_count = [0]

                        async def scan_one(idx: int, data: dict):
                            if self._stop.is_set():
                                return
                            async with sem:
                                ig = await find_instagram_on_site(ctx, data["website"])
                                done_count[0] += 1
                                q.put({"type": "progress",
                                       "current": done_count[0], "total": len(to_scan)})
                                if ig:
                                    data["instagram"] = ig
                                    database.update_instagram(data["maps_url"], ig)
                                    q.put({"type": "ig_update", "idx": idx, "instagram": ig})

                        await asyncio.gather(*[scan_one(i, d) for i, d in to_scan])

                # ── Salvar histórico no banco ──────────────────────────────
                database.save_search(
                    query=query, location=location,
                    total_urls=len(urls),
                    total_new=new_count,
                    total_cached=cached_count,
                )

                await ctx.close()
                await browser.close()
                q.put({"type": "done"})

        except Exception as exc:
            q.put({"type": "error", "text": str(exc)})

    # ── Banco de dados ────────────────────────────────────────────────────────

    def _refresh_db_stats(self):
        stats = database.get_stats()
        self._lbl_db_stats.configure(
            text=(
                f"{stats['total']:,} empresas salvas\n"
                f"{stats['with_website']:,} com site  ·  "
                f"{stats['with_instagram']:,} com Instagram"
            )
        )

    def _show_history(self):
        win = ctk.CTkToplevel(self)
        win.title("Histórico de Buscas")
        win.geometry("780x480")
        win.resizable(True, True)
        win.grab_set()
        win.configure(fg_color=CLR_BG3)

        # Stats no topo
        stats = database.get_stats()
        ctk.CTkLabel(
            win,
            text=(
                f"Banco  —  {stats['total']:,} empresas  ·  "
                f"{stats['with_website']:,} com site  ·  "
                f"{stats['with_instagram']:,} com Instagram  ·  "
                f"{stats['searches']:,} buscas realizadas"
            ),
            font=ctk.CTkFont(size=12),
            text_color=CLR_MUTED,
        ).pack(padx=16, pady=(16, 8), anchor="w")

        # Tabela de histórico
        frame = ctk.CTkFrame(win, corner_radius=8, fg_color=CLR_BG2)
        frame.pack(fill="both", expand=True, padx=16, pady=(0, 16))
        frame.grid_rowconfigure(0, weight=1)
        frame.grid_columnconfigure(0, weight=1)

        cols   = ("searched_at", "query", "location", "total_urls", "total_new", "total_cached")
        labels = ("Data / Hora",  "Nicho", "Região",   "URLs",       "Novas",     "Cache")
        widths = (135, 210, 130, 70, 70, 70)
        anchors = ("center", "w", "w", "center", "center", "center")

        tree = ttk.Treeview(
            frame, columns=cols, show="headings",
            style="Maps.Treeview", selectmode="browse",
        )
        for col, lbl, w, anch in zip(cols, labels, widths, anchors):
            tree.heading(col, text=lbl)
            tree.column(col, width=w, anchor=anch)

        vsb = ttk.Scrollbar(frame, orient="vertical", command=tree.yview)
        tree.configure(yscrollcommand=vsb.set)
        tree.grid(row=0, column=0, sticky="nsew", padx=(8, 0), pady=8)
        vsb.grid(row=0, column=1, sticky="ns", pady=8)

        tree.tag_configure("odd",  background="#1e293b")
        tree.tag_configure("even", background="#253047")

        for i, s in enumerate(database.get_searches(100)):
            tag = "even" if i % 2 == 0 else "odd"
            tree.insert("", "end", tags=(tag,), values=(
                s["searched_at"], s["query"], s["location"],
                s["total_urls"], s["total_new"], s["total_cached"],
            ))

        # Botão limpar banco
        btn_frame = ctk.CTkFrame(win, fg_color="transparent")
        btn_frame.pack(fill="x", padx=16, pady=(0, 12))

        def confirm_clear():
            if messagebox.askyesno(
                "Limpar banco",
                "Isso vai apagar TODOS os dados salvos.\nTem certeza?",
                parent=win,
            ):
                database.clear_all()
                self._refresh_db_stats()
                for item in tree.get_children():
                    tree.delete(item)
                messagebox.showinfo("Pronto", "Banco limpo com sucesso.", parent=win)

        ctk.CTkButton(
            btn_frame, text="Limpar banco de dados",
            height=32, font=ctk.CTkFont(size=11),
            fg_color="#7f1d1d", hover_color=CLR_RED,
            command=confirm_clear,
        ).pack(side="right")

    # ── Exportação ────────────────────────────────────────────────────────────

    def _export_xlsx(self):
        if not self._results:
            messagebox.showwarning("Atenção", "Nenhum resultado para exportar.")
            return

        path = filedialog.asksaveasfilename(
            defaultextension=".xlsx",
            filetypes=[("Excel", "*.xlsx")],
            initialfile=f"maps_{datetime.now().strftime('%Y%m%d_%H%M')}.xlsx",
            title="Salvar planilha Excel",
        )
        if not path:
            return

        try:
            import openpyxl
            from openpyxl.styles import (
                Alignment, Font, GradientFill, PatternFill, Side, Border,
            )
            from openpyxl.utils import get_column_letter

            wb = openpyxl.Workbook()
            ws = wb.active
            ws.title = "Maps Scraper"

            # Cabeçalho
            hdr_fill = PatternFill("solid", fgColor="1E3A5F")
            hdr_font = Font(color="FFFFFF", bold=True, size=11, name="Segoe UI")
            hdr_align = Alignment(horizontal="center", vertical="center", wrap_text=False)

            for ci, header in enumerate(CSV_HEADERS, 1):
                cell = ws.cell(row=1, column=ci, value=header)
                cell.fill  = hdr_fill
                cell.font  = hdr_font
                cell.alignment = hdr_align

            ws.row_dimensions[1].height = 24

            # Estilos de linha
            fill_even = PatternFill("solid", fgColor="EFF6FF")
            font_link = Font(color="1D4ED8", underline="single", name="Segoe UI", size=10)
            font_ig   = Font(color="9D174D", underline="single", name="Segoe UI", size=10)
            font_base = Font(name="Segoe UI", size=10)
            align_l   = Alignment(horizontal="left",   vertical="center")
            align_c   = Alignment(horizontal="center", vertical="center")

            for ri, data in enumerate(self._results, 2):
                fill = fill_even if ri % 2 == 0 else None

                for ci, key in enumerate(CSV_KEYS, 1):
                    val  = data.get(key, "")
                    cell = ws.cell(row=ri, column=ci, value=val)
                    cell.font = font_base

                    if fill:
                        cell.fill = fill

                    if key == "website" and val and val.startswith("http"):
                        cell.hyperlink = val
                        cell.font = font_link
                        cell.alignment = align_l
                    elif key == "instagram" and val:
                        url = f"https://{val}" if not val.startswith("http") else val
                        cell.hyperlink = url
                        cell.font = font_ig
                        cell.alignment = align_l
                    elif key == "maps_url" and val:
                        cell.hyperlink = val
                        cell.font = font_link
                        cell.alignment = align_l
                    elif key in ("avaliacao", "total_avaliacoes"):
                        cell.alignment = align_c
                    else:
                        cell.alignment = align_l

                ws.row_dimensions[ri].height = 20

            # Larguras automáticas
            for ci, key in enumerate(CSV_KEYS, 1):
                col_letter = get_column_letter(ci)
                max_len    = len(CSV_HEADERS[ci - 1])
                for data in self._results:
                    val = str(data.get(key, ""))
                    if len(val) > max_len:
                        max_len = len(val)
                ws.column_dimensions[col_letter].width = min(max_len + 3, 55)

            # Congelar cabeçalho + auto-filtro
            ws.freeze_panes = "A2"
            ws.auto_filter.ref = ws.dimensions

            wb.save(path)
            messagebox.showinfo("Exportado!", f"Planilha salva em:\n{path}")

        except ImportError:
            messagebox.showerror(
                "Dependência faltando",
                "Instale openpyxl para exportar Excel:\n\npip install openpyxl",
            )
        except Exception as exc:
            messagebox.showerror("Erro ao exportar", str(exc))

    def _export_csv(self):
        if not self._results:
            messagebox.showwarning("Atenção", "Nenhum resultado para exportar.")
            return

        path = filedialog.asksaveasfilename(
            defaultextension=".csv",
            filetypes=[("CSV", "*.csv")],
            initialfile=f"maps_{datetime.now().strftime('%Y%m%d_%H%M')}.csv",
            title="Salvar CSV",
        )
        if not path:
            return

        try:
            with open(path, "w", newline="", encoding="utf-8-sig") as f:
                w = csv.writer(f)
                w.writerow(CSV_HEADERS)
                for data in self._results:
                    w.writerow([data.get(k, "") for k in CSV_KEYS])
            messagebox.showinfo("Exportado!", f"CSV salvo em:\n{path}")
        except Exception as exc:
            messagebox.showerror("Erro ao exportar", str(exc))


# ── Entrada ───────────────────────────────────────────────────────────────────

if __name__ == "__main__":
    app = App()
    app.mainloop()
